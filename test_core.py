import unittest, tempfile, io
from pathlib import Path
from unittest.mock import patch
from datetime import date
from decimal import Decimal
import openpyxl
from core import *
from exports import make_excel, make_pdf, summary_cards

HEADERS=[KEY,"Megbízó","Alvállalkozó",CONTACT,"Alvállalkozói díj","Alvállalkozói díj pénzneme","Díj","Díj pénzneme",
         MARGIN,CURRENCY,"Árrés (állapot szerint)","Árrés (állapot szerint) pénzneme","Számlázás","Számla sorszáma",
         "Teljesítettség állapota",DATE]
def rec(key="1-1",margin="100",currency="EUR",dates="2026-07-15",owner="Peti"):
    original=dict(zip(HEADERS,[key,"Megbízó Kft.","Alvállalkozó Kft.",owner,10,currency,float(dec(margin)+10),currency,
                              margin,currency,1,"EUR","26/123","26/123","Teljesítve",dates]))
    return Record(key,owner,original,dec(margin),currency,parse_dates(dates),fingerprint(original))
def fx(month):
    return {"rate":"400","source":"teszt","method":"teszt","count":22,"daily":{}}

class Tests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.s=Store(Path(self.tmp.name)/"db.sqlite3")
        self.s.set_setting("start_month","2026-07")
    def tearDown(self):self.tmp.cleanup()
    def plan(self,r,month="2026-07",owner="Peti"):
        return build_plan(r,HEADERS,self.s,month,owner,"hash",rate_provider=fx)
    def test_same_month_dates_and_negative(self):
        a=rec(dates="2026-07-01, 2026-07-31")
        b=rec("2",margin="-10",currency="HUF")
        p=self.plan([a,b])
        self.assertEqual(p["total"],4000)
        self.assertFalse(p["blockers"]);self.assertTrue(p["lines"][1]["negative"])
    def test_cross_month_requires_choice_and_retains_assignment(self):
        a=rec(dates="2026-07-01, 2026-08-05")
        self.assertTrue(self.plan([a],"2026-08")["blockers"])
        self.s.save_decision(a,"2026-07","Ellenőrizve")
        self.assertFalse(self.plan([a])["blockers"])
        a.original[DATE]="2026-07-01, 2026-09-01";a.dates=parse_dates(a.original[DATE]);a.signature=fingerprint(a.original)
        self.assertFalse(self.plan([a],"2026-09")["blockers"])
        self.assertEqual(origin(a,self.s),"2026-07")
    def test_repeat_close_and_late_and_changed(self):
        a=rec()
        p=self.plan([a]);self.s.close(p)
        with self.assertRaises(ValueError):self.s.close(p)
        b=rec("2",margin="10",currency="HUF")
        p=self.plan([rec(margin="150"),b],"2026-08")
        self.assertEqual(p["total"],2001);self.assertEqual(p["current"],0)
        self.s.close(p)
        self.assertEqual(self.plan([rec(margin="150"),b],"2026-09")["total"],0)
    def test_stale_revision_and_missing(self):
        p=self.plan([rec()]);self.s.set_setting("test","x")
        with self.assertRaises(ValueError):self.s.close(p)
        self.s.close(self.plan([rec()]))
        p=self.plan([],"2026-08");self.assertTrue(p["warnings"]);self.assertEqual(p["total"],0)
    def test_contact_move(self):
        self.s.close(self.plan([rec()]))
        moved=rec(owner="Anna")
        p=self.plan([moved],"2026-08")
        self.assertEqual(p["total"],-4000);self.s.close(p)
        p=self.plan([moved],"2026-08","Anna")
        self.assertEqual(p["total"],4000);self.s.close(p)
        self.assertEqual(self.plan([moved],"2026-09","Anna")["total"],0)
    def test_rounding_difference(self):
        self.s.close(self.plan([rec(margin="15",currency="HUF")]))
        p=self.plan([rec(margin="16",currency="HUF")],"2026-08")
        self.assertEqual(p["total"],0);self.assertEqual(len(p["lines"]),1)
    def test_history_frozen(self):
        p=self.plan([rec()]);cid=self.s.close(p)
        self.assertEqual(self.s.closing(cid)["total"],4000)
        self.s.save_rate("2026-07",dict(fx(""),rate="500"))
        self.assertEqual(self.s.rate("2026-07")["rate"],"400")
        self.assertEqual(self.s.closing(cid)["lines"][0]["commission"],4000)
    def test_parser_duplicate_missing_and_malformed(self):
        w=openpyxl.Workbook();s=w.active;s.append(HEADERS)
        s.append(list(rec().original.values()));s.append(list(rec().original.values()))
        out=io.BytesIO();w.save(out)
        h,rows,errors=load_workbook(out.getvalue())
        self.assertEqual(len(rows),1);self.assertEqual(len(errors),1)
        with self.assertRaises(ValueError):parse_dates("2026-07-01 + hibás")
        with self.assertRaises(ValueError):parse_dates("2026-02-30")
        with self.assertRaises(ValueError):dec(None)
    def test_exports_preserve_columns_and_total(self):
        p=self.plan([rec(),rec("2",margin="-10",currency="HUF")])
        x=make_excel(p);w=openpyxl.load_workbook(io.BytesIO(x),data_only=True)
        s=w["Jutalékelszámolás"]
        self.assertEqual([c.value for c in s[8]],export_columns(p))
        self.assertEqual(s.cell(14,len(export_columns(p))).value,4000)
        self.assertEqual(len(s._images),1)
        pdf=make_pdf(p)
        self.assertTrue(pdf.startswith(b"%PDF"))
        from pypdf import PdfReader
        d=PdfReader(io.BytesIO(pdf))
        self.assertAlmostEqual(float(d.pages[0].mediabox.width),841.8898,places=2)
        self.assertAlmostEqual(float(d.pages[0].mediabox.height),595.2756,places=2)
        text="".join(page.extract_text() for page in d.pages)
        self.assertIn("4 000",text)
        for h in export_columns(p):self.assertIn("".join(h.split()),"".join(text.split()))
        self.assertEqual(len(export_columns(p)),9)
        self.assertNotIn("Előzőleg elszámolt jutalék (Ft)",export_columns(p))
        self.assertIn("100,00 EUR",text)
        self.assertEqual(str(s.page_setup.paperSize),str(s.PAPERSIZE_A4))
        # Az Excel a pénzösszegeket számként őrzi meg, a pénznem számformátum.
        self.assertIsInstance(s.cell(9,6).value,(int,float))
        self.assertIn("EUR",s.cell(9,6).number_format)
        self.assertNotIn("Árrés",export_columns(p))
        self.assertIn("Árrés\nextrákkal",export_columns(p))
        self.assertIn("2026.07.15",d.pages[1].extract_text())
        self.assertEqual(s.cell(9,8).value,"2026.07.15")
        self.assertEqual(w.sheetnames[0],"Összesítő")
        self.assertIn("ÖSSZESÍTŐ",d.pages[0].extract_text())
        self.assertNotIn("RÉSZLETEZŐ",d.pages[0].extract_text())
        self.assertIn("RÉSZLETEZŐ",d.pages[1].extract_text())
        self.assertIn("4 000 Ft",d.pages[1].extract_text())
        for col in range(4,10):
            self.assertEqual(s.cell(9,col).alignment.horizontal,"center")
            self.assertEqual(s.cell(8,col).alignment.horizontal,"center")
            self.assertEqual(s.cell(10,col).font.color.rgb,"009C2222")

    def test_snapshot_repeats_ignores_closings_and_other_months(self):
        a=rec()
        self.s.close(self.plan([a]))
        self.s.set_setting("start_month","2026-09")
        for _ in range(2):
            p=build_plan([a,rec("later",dates="2026-08-03")],HEADERS,self.s,"2026-07","Peti","x",rate_provider=fx,snapshot=True)
            self.assertEqual(p["total"],4000)
            self.assertEqual([x["key"] for x in p["lines"]],[a.key])
            self.assertEqual(p["status"],"AKTUÁLIS ÁLLAPOT")
            self.assertEqual(p["corrections"],0)

    def test_snapshot_without_start_and_changed_margin(self):
        s=Store(Path(self.tmp.name)/"new.sqlite3")
        p=build_plan([rec(margin="-20",currency="HUF")],HEADERS,s,"2026-07","Peti","x",rate_provider=fx,snapshot=True)
        self.assertEqual(p["total"],0)
        p=build_plan([rec(margin="80",currency="HUF")],HEADERS,s,"2026-07","Peti","x",rate_provider=fx,snapshot=True)
        self.assertEqual(p["total"],8)

    def test_negative_cards_only_for_selected_month(self):
        records=[rec("negative",margin="-10",dates="2026-07-15"),rec("positive",dates="2026-08-15")]
        july=build_plan(records,HEADERS,self.s,"2026-07","Peti","x",rate_provider=fx,snapshot=True)
        august=build_plan(records,HEADERS,self.s,"2026-08","Peti","x",rate_provider=fx,snapshot=True)
        self.assertEqual(len(summary_cards(july)),6)
        self.assertEqual(len(summary_cards(august)),4)
        self.assertFalse(any("Negatív" in label for label,value,unit,tone in summary_cards(august)))
        from pypdf import PdfReader
        text=PdfReader(io.BytesIO(make_pdf(august))).pages[0].extract_text()
        self.assertNotIn("Negatív árrésű fuvarok",text)
        self.assertNotIn("Negatív tételek jutaléka",text)
        w=openpyxl.load_workbook(io.BytesIO(make_excel(august)))
        self.assertFalse(any("Negatív" in str(c.value) for row in w["Összesítő"] for c in row))

    def test_snapshot_cross_month_checks_only_relevant_month(self):
        a=rec(dates="2026-07-01, 2026-08-05")
        p=build_plan([a],HEADERS,self.s,"2026-09","Peti","x",rate_provider=fx,snapshot=True)
        self.assertFalse(p["blockers"])
        p=build_plan([a],HEADERS,self.s,"2026-08","Peti","x",rate_provider=fx,snapshot=True)
        self.assertTrue(p["blockers"])
        self.s.save_decision(a,"2026-08","Ellenőrizve")
        p=build_plan([a],HEADERS,self.s,"2026-08","Peti","x",rate_provider=fx,snapshot=True)
        self.assertEqual(p["total"],4000)
    def test_real_source(self):
        source=Path(__file__).parents[1]/"commission_review"/"source.xlsx"
        if not source.exists():self.skipTest("Forrás csak fejlesztői környezetben")
        h,rows,errors=load_workbook(source.read_bytes())
        self.assertEqual(len(rows),2611);self.assertEqual(errors,[])
        self.assertEqual(sum(r.margin<0 for r in rows),2)
        self.assertEqual(sum(len(r.months)>1 for r in rows),419)
        self.s.set_setting("start_month","2026-09")
        p=build_plan(rows,h,self.s,"2026-09","Buday András","real",rate_provider=fx)
        self.assertFalse(p["blockers"])
        self.assertGreater(len(p["lines"]),0)

    def test_failed_rate_requested_once(self):
        calls=[]
        def fail(month):
            calls.append(month)
            raise ValueError("Nincs kapcsolat")
        p=build_plan([rec(),rec("2")],HEADERS,self.s,"2026-07","Peti","x",rate_provider=fail)
        self.assertEqual(len(calls),1)
        self.assertEqual(len(p["blockers"]),1)
        self.assertEqual(p["lines"],[])
    def test_change_during_calculation_blocks_close(self):
        def altered(month):
            self.s.set_setting("concurrent","changed")
            return fx(month)
        p=build_plan([rec()],HEADERS,self.s,"2026-07","Peti","x",rate_provider=altered)
        self.assertTrue(p["blockers"])
        with self.assertRaises(ValueError):self.s.close(p)
    def test_backup_reopens_with_ledger(self):
        self.s.close(self.plan([rec()]))
        restored=Path(self.tmp.name)/"restored.sqlite3"
        restored.write_bytes(self.s.backup())
        s=Store(restored)
        self.assertEqual(len(s.closings()),1)
        self.assertEqual(s.latest()[("1-1","Peti")]["commission"],4000)

if __name__=="__main__":unittest.main(verbosity=2)
