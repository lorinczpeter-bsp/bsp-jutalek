"""Átsorolások, kifizetési bizonyítékok és elemzés közös üzleti szabályai."""
import io,re,hashlib,json
from pathlib import Path
from datetime import datetime
from decimal import Decimal
import openpyxl
from core import dec,valid_month,KEY

def init_registry(store):
    with store.connect() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS overrides(freight TEXT PRIMARY KEY,month TEXT NOT NULL,note TEXT,created TEXT);
        CREATE TABLE IF NOT EXISTS payments(freight TEXT PRIMARY KEY,contact TEXT NOT NULL,amount TEXT NOT NULL,
            paid_date TEXT NOT NULL,source TEXT NOT NULL,source_hash TEXT NOT NULL,created TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS registry_audit(id INTEGER PRIMARY KEY AUTOINCREMENT,action TEXT,freight TEXT,payload TEXT,created TEXT);
        ''')

def overrides(store):
    init_registry(store)
    with store.connect() as c:return {r[0]:r[1] for r in c.execute('SELECT freight,month FROM overrides')}

def resolved_months(records,store):
    with store.connect() as c:
        moved={r[0]:r[1] for r in c.execute('SELECT freight,month FROM overrides')}
        decisions={r[0]:(r[1],r[2]) for r in c.execute('SELECT freight,signature,month FROM decisions')}
    result={}
    for r in records:
        d=decisions.get(r.key)
        result[r.key]=moved.get(r.key) or (r.months[0] if len(r.months)==1 else (d[1] if d and d[0]==r.signature else None))
    return result

def reassign(store,record,month,note=''):
    if not valid_month(month):raise ValueError('Érvényes ÉÉÉÉ-HH hónap szükséges.')
    init_registry(store)
    with store.connect() as c:
        old=c.execute('SELECT month FROM overrides WHERE freight=?',(record.key,)).fetchone()
        c.execute('INSERT OR REPLACE INTO overrides VALUES(?,?,?,?)',(record.key,month,note.strip(),datetime.now().isoformat()))
        c.execute('INSERT INTO registry_audit(action,freight,payload,created) VALUES(?,?,?,?)',
            ('reassign',record.key,json.dumps({'previous':old[0] if old else None,'month':month,'note':note},ensure_ascii=False),datetime.now().isoformat()))
        store.bump(c)
    if overrides(store).get(record.key)!=month:raise ValueError('Az átsorolás visszaellenőrzése sikertelen.')

def clear_reassignment(store,key):
    init_registry(store)
    with store.connect() as c:
        c.execute('DELETE FROM overrides WHERE freight=?',(key,))
        c.execute('INSERT INTO registry_audit(action,freight,payload,created) VALUES(?,?,?,?)',('clear_reassignment',key,'{}',datetime.now().isoformat()))
        store.bump(c)

def payments(store):
    init_registry(store)
    with store.connect() as c:
        c.row_factory=__import__('sqlite3').Row
        return {r['freight']:dict(r) for r in c.execute('SELECT * FROM payments ORDER BY paid_date,freight')}

def read_payment_report(data,contact=''):
    """A kézi és az appos XLSX tényleges Jutalék oszlopát olvassa, nem a kalkulált auditösszeget."""
    wb=openpyxl.load_workbook(io.BytesIO(data),data_only=True,read_only=True)
    found=[]
    inferred=contact.strip()
    if 'Összesítő' in wb.sheetnames:
        for row in wb['Összesítő'].iter_rows(min_row=2,max_row=2,values_only=True):
            names=[str(v).strip() for v in row if v]
            if len(names)==1:inferred=names[0]
    for sheet in wb:
        for idx,row in enumerate(sheet.iter_rows(min_row=1,max_row=15,values_only=True),1):
            labels=[re.sub(r'\s+',' ',str(v or '')).strip() for v in row]
            if KEY in labels:
                commissions=[i for i,s in enumerate(labels) if s in ['Jutalék','Jutalék (Ft)','Jutalék – 10% (Ft)']]
                if len(commissions)==1:found.append((sheet,idx,labels.index(KEY),commissions[0]))
    if len(found)!=1:
        wb.close();raise ValueError('Pontosan egy Fuvarszám és Jutalék oszlopot tartalmazó részletező szükséges.')
    sheet,header,key_col,amount_col=found[0]
    rows=[];errors=[];seen=set()
    for n,row in enumerate(sheet.iter_rows(min_row=header+1,values_only=True),header+1):
        key=str(row[key_col] or '').strip()
        if not re.fullmatch(r'\d+-\d+',key):continue
        try:
            if key in seen:raise ValueError('Ismétlődő fuvarszám a korábbi kimutatásban.')
            seen.add(key)
            if amount_col>=len(row) or row[amount_col] is None:raise ValueError('Hiányzó jutalék vagy Excel-képlet eltárolt eredmény nélkül.')
            raw=row[amount_col]
            if isinstance(raw,str):raw=re.sub(r'\s*(Ft|HUF)\s*$','',raw,flags=re.I)
            amount=dec(raw)
            if amount<0:raise ValueError('Negatív kifizetés: kézi egyeztetés szükséges.')
            if amount==0:continue
            if not inferred:raise ValueError('A fuvarszervező nevét meg kell adni.')
            rows.append({'freight':key,'contact':inferred,'amount':str(amount),'row':n})
        except ValueError as e:errors.append({'Sor':n,'Fuvarszám':key,'Hiba':str(e)})
    wb.close()
    if not rows and not errors:raise ValueError('Nincs pozitív jutalékot tartalmazó fuvar a kimutatásban.')
    return rows,errors

def import_payments(store,rows,source,source_hash,paid_date):
    if not rows:raise ValueError('Nincs importálható kifizetés.')
    __import__('datetime').date.fromisoformat(str(paid_date))
    if any(dec(r['amount'])<=0 or not r['contact'].strip() for r in rows):raise ValueError('Hibás kifizetési adatok.')
    if len({r['freight'] for r in rows})!=len(rows):raise ValueError('Ismétlődő fuvarszám.')
    init_registry(store);new=0;duplicates=[]
    with store.connect() as c:
        c.execute('BEGIN IMMEDIATE')
        for r in rows:
            if c.execute('SELECT 1 FROM payments WHERE freight=?',(r['freight'],)).fetchone():
                duplicates.append(r['freight']);continue
            c.execute('INSERT INTO payments VALUES(?,?,?,?,?,?,?)',(r['freight'],r['contact'],r['amount'],str(paid_date),source,source_hash,datetime.now().isoformat()))
            c.execute('INSERT INTO registry_audit(action,freight,payload,created) VALUES(?,?,?,?)',('payment_import',r['freight'],json.dumps(r,ensure_ascii=False),datetime.now().isoformat()))
            new+=1
        if new:store.bump(c)
    return new,duplicates

def remove_payment(store,key):
    init_registry(store)
    with store.connect() as c:
        old=c.execute('SELECT * FROM payments WHERE freight=?',(key,)).fetchone()
        if not old:raise ValueError('Nincs ilyen nyilvántartott kifizetés.')
        c.execute('INSERT INTO registry_audit(action,freight,payload,created) VALUES(?,?,?,?)',('payment_remove',key,json.dumps(old,ensure_ascii=False),datetime.now().isoformat()))
        c.execute('DELETE FROM payments WHERE freight=?',(key,));store.bump(c)

def filename_stem(plan):
    slug=''.join(c if c.isalnum() or c in '-_' else '_' for c in plan['contact'])
    return f"Jutalek_{plan['month']}_{slug}_{__import__('core').today().isoformat()}"

def analytics(records,headers,store,months,contacts,rate_provider=None):
    from core import build_plan
    rows=[];issues=[]
    for month in sorted(set(months)):
        for contact in sorted(set(contacts)):
            p=build_plan(records,headers,store,month,contact,'analytics',snapshot=True,rate_provider=rate_provider)
            issues.extend(f'{month} / {contact}: {msg}' for msg in p['blockers'])
            rows.append({'Hónap':month,'Fuvarszervező':contact,'Számított jutalék (Ft)':p['earned_total'],
                'Kifizetettként kezelt (Ft)':p['paid_total'],'Még fizetendő (Ft)':p['total'],'Fuvarok':len(p['lines']),
                'Hiányos':bool(p['blockers'])})
    return rows,list(dict.fromkeys(issues))
