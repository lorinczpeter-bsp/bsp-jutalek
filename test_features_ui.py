import io,tempfile,unittest
from pathlib import Path
import openpyxl
from core import Store,origin,load_workbook
from settlement import payments
from ui_features import show_analytics,show_payments,show_reassign
from test_ui import App,Rerun
from test_core import HEADERS,rec

class FeatureApp(App):
    def __init__(self):
        super().__init__(b'');self.submit=False;self.clicked='';self.confirm=False;self.report=None;self.tables=[];self.metrics=[];self.charts=0
    def text_input(self,label,*a,**kw):return '1-1' if label=='Átsorolandó fuvarszám' else kw.get('value','')
    def selectbox(self,label,options,index=0,**kw):return '2026-08' if label=='Új elszámolási hónap' else options[index]
    def multiselect(self,label,options,default=None,**kw):return default
    def metric(self,label,value,*a,**kw):self.metrics.append((label,value))
    def dataframe(self,data,**kw):self.tables.append(data)
    def bar_chart(self,*a,**kw):self.charts+=1
    def line_chart(self,*a,**kw):self.charts+=1
    def vega_lite_chart(self,*a,**kw):self.charts+=1
    def file_uploader(self,*a,**kw):return self.report
    def date_input(self,*a,**kw):return kw['value']
    def checkbox(self,*a,**kw):return self.confirm
    def button(self,label,**kw):return label==self.clicked and not kw.get('disabled',False)
    def form_submit_button(self,label,*a,**kw):return self.submit and label=='Átsorolás mentése'

class FeatureUITests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.s=Store(Path(self.tmp.name)/'db.sqlite3');self.ui=FeatureApp()
        self.records=[rec('1-1',currency='HUF',margin='100',dates='2026-07-01'),rec('2-1',currency='HUF',margin='200',dates='2026-09-01',owner='Anna')]
        self.bundle=('x',HEADERS,self.records,[])
    def tearDown(self):self.tmp.cleanup()
    def test_analytics_multiple_people_months_and_empty_month(self):
        show_analytics(self.ui,self.s,self.bundle)
        self.assertEqual(self.ui.charts,2)
        self.assertIn(('Kiválasztott hónapok összege','30 Ft'),self.ui.metrics)
        self.assertIn(('Havi átlag','10,00 Ft'),self.ui.metrics)
        self.assertEqual(len(self.ui.tables[0]),3)
    def test_reassign_form_saves_requested_month(self):
        self.ui.submit=True
        with self.assertRaises(Rerun):show_reassign(self.ui,self.s,self.bundle)
        self.assertEqual(origin(self.records[0],self.s),'2026-08')
    def test_payment_upload_needs_confirmation_and_saves(self):
        w=openpyxl.Workbook();w.active.append(['Fuvarszám','Jutalék']);w.active.append(['1-1',10])
        out=io.BytesIO();w.save(out)
        report=io.BytesIO(out.getvalue());report.name='paid.xlsx';self.ui.report=report
        self.ui.clicked='Kifizetési előzmények rögzítése'
        show_payments(self.ui,self.s,self.bundle)
        self.assertFalse(payments(self.s))
        self.ui.confirm=True
        with self.assertRaises(Rerun):show_payments(self.ui,self.s,self.bundle)
        self.assertEqual(len(payments(self.s)),1)

if __name__=='__main__':unittest.main()
