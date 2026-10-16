import io,tempfile,unittest,hashlib
from pathlib import Path
from decimal import Decimal
from unittest.mock import patch
import openpyxl
from core import Store,build_plan,origin,MARGIN,CURRENCY,fingerprint,Record,dec,parse_dates
from test_core import HEADERS,rec
from exports import make_excel,make_pdf
from settlement import *

RATE='365.0531818181818181818181818'
def fx(month):return {'rate':RATE if month=='2026-09' else '400.12789','source':'teszt','daily':{},'count':22}
def mixed(key='16691-1',reverse=False):
    r=rec(key,dates='2026-09-24')
    r.original.update({'Díj':85000 if reverse else 420,'Díj pénzneme':'HUF' if reverse else 'EUR',
        'Alvállalkozói díj':200 if reverse else 85000,'Alvállalkozói díj pénzneme':'EUR' if reverse else 'HUF'})
    return r

class SettlementTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.s=Store(Path(self.tmp.name)/'db.sqlite3')
    def tearDown(self):self.tmp.cleanup()
    def plan(self,rows,month='2026-09',contact='Peti'):
        return build_plan(rows,HEADERS,self.s,month,contact,'test',rate_provider=fx,snapshot=True)
    def test_mixed_uses_full_precision_and_exports_huf(self):
        p=self.plan([mixed()]);line=p['lines'][0]
        expected=Decimal('420')*Decimal(RATE)-Decimal('85000')
        self.assertEqual(dec(line['margin_huf']),expected)
        self.assertEqual(line['commission'],6832)
        self.assertEqual(line['original'][CURRENCY],'HUF')
        self.assertEqual(line['source_original'][MARGIN],'100')
        self.assertEqual(self.s.rate('2026-09')['rate'],RATE)
        w=openpyxl.load_workbook(io.BytesIO(make_excel(p)),data_only=True)
        self.assertIn('365,05 Ft/EUR',w['Jutalékelszámolás']['A5'].value)
        self.assertIn('HUF',w['Jutalékelszámolás']['F9'].number_format)
        self.assertEqual(w['Árfolyamnapló']['B2'].number_format,'0.00')
    def test_reverse_mixed_and_new_negative_sign(self):
        r=mixed(reverse=True)
        self.assertEqual(self.plan([r])['total'],1199)
        r.original['Díj']=100
        p=self.plan([r]);self.assertTrue(p['lines'][0]['negative']);self.assertEqual(p['total'],0)
    def test_mixed_missing_fee_is_blocker(self):
        r=mixed();r.original['Díj']=None
        p=self.plan([r]);self.assertTrue(p['blockers']);self.assertEqual(p['lines'],[])
    def test_same_currency_keeps_export_margin(self):
        r=rec(margin='123',dates='2026-09-01');r.original['Díj']=999999
        self.assertEqual(self.plan([r])['total'],4490)
    def test_reassignment_any_month_moves_once_and_changes_fx(self):
        r=mixed();reassign(self.s,r,'2026-08','Visszamenőleges rendezés')
        self.assertEqual(origin(r,self.s),'2026-08')
        self.assertEqual(self.plan([r])['lines'],[])
        august=self.plan([r],'2026-08')
        self.assertEqual(dec(august['lines'][0]['margin_huf']),Decimal(420)*Decimal('400.12789')-85000)
        self.assertEqual(len(august['lines']),1)
        clear_reassignment(self.s,r.key)
        self.assertEqual(origin(r,self.s),'2026-09')
        with self.assertRaises(ValueError):reassign(self.s,r,'2026-13')
    def test_report_import_old_excel_and_duplicate_protection(self):
        w=openpyxl.Workbook();ws=w.active;ws.append(['Fuvarszám','Jutalék']);ws.append(['16691-1',6832.1]);ws.append(['16692-1',0])
        out=io.BytesIO();w.save(out)
        rows,errors=read_payment_report(out.getvalue(),'Peti');self.assertFalse(errors);self.assertEqual(len(rows),1)
        self.assertEqual(import_payments(self.s,rows,'kézi.xlsx','hash','2026-10-10'),(1,[]))
        self.assertEqual(import_payments(self.s,rows,'másik.xlsx','hash2','2026-10-11'),(0,['16691-1']))
        p=self.plan([mixed()]);self.assertEqual(p['total'],0);self.assertEqual(p['earned_total'],6832)
        self.assertEqual(p['paid_keys'],['16691-1']);self.assertEqual(p['actual_paid_total'],'6832.1')
        reassign(self.s,mixed(),'2026-08')
        self.assertEqual(self.plan([mixed()],'2026-08')['total'],0)
        remove_payment(self.s,'16691-1')
        self.assertGreater(self.plan([mixed()],'2026-08')['total'],0)
    def test_report_new_format_infers_owner_and_skips_paid_zero(self):
        p=self.plan([mixed()]);rows,errors=read_payment_report(make_excel(p))
        self.assertFalse(errors);self.assertEqual(rows[0]['contact'],'Peti')
        import_payments(self.s,rows,'új.xlsx','x','2026-10-10')
        p=self.plan([mixed()]);w=openpyxl.load_workbook(io.BytesIO(make_excel(p)))
        self.assertIn('Kifizetési ellenőrzés',w.sheetnames)
        self.assertIn('kifizetve',w['Jutalékelszámolás']['I9'].number_format)
        with self.assertRaises(ValueError):read_payment_report(make_excel(p))
    def test_report_duplicates_missing_cached_formula_negative(self):
        for values in [[['16691-1',10],['16691-1',10]],[['16691-1','=1+1']],[['16691-1',-10]]]:
            w=openpyxl.Workbook();w.active.append(['Fuvarszám','Jutalék'])
            for row in values:w.active.append(row)
            out=io.BytesIO();w.save(out)
            rows,errors=read_payment_report(out.getvalue(),'Peti');self.assertTrue(errors)
    def test_payment_owner_change_does_not_pay_twice(self):
        import_payments(self.s,[{'freight':'16691-1','contact':'Másik név','amount':'1000'}],'a.xlsx','x','2026-10-10')
        p=self.plan([mixed()]);self.assertEqual(p['total'],0);self.assertTrue(any('eltér' in x for x in p['warnings']))
    def test_analytics_zero_months_multiple_contacts_and_paid(self):
        records=[rec('1-1',margin='100',currency='HUF',dates='2026-07-01'),
                 rec('2-1',margin='200',currency='HUF',dates='2026-09-01',owner='Anna')]
        import_payments(self.s,[{'freight':'1-1','contact':'Peti','amount':'10'}],'a.xlsx','x','2026-08-01')
        rows,issues=analytics(records,HEADERS,self.s,['2026-07','2026-08','2026-09'],['Peti','Anna'],fx)
        self.assertFalse(issues);self.assertEqual(len(rows),6)
        self.assertEqual(sum(r['Számított jutalék (Ft)'] for r in rows),30)
        self.assertEqual(sum(r['Még fizetendő (Ft)'] for r in rows),20)
        self.assertEqual(sum(r['Kifizetettként kezelt (Ft)'] for r in rows),10)
    def test_filename_uses_creation_date_not_status(self):
        with patch('core.today',return_value=__import__('datetime').date(2026,10,17)):
            name=filename_stem({'contact':'Anikó','month':'2026-09','created':'2020-01-01','status':'AKTUÁLIS ÁLLAPOT'})
        self.assertEqual(name,'Jutalek_2026-09_Anikó_2026-10-17')

if __name__=='__main__':unittest.main()
