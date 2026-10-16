import io,tempfile,unittest
from pathlib import Path
from decimal import Decimal
import openpyxl
from pypdf import PdfReader
from core import Store,build_plan,export_rows,MARGIN
from exports import make_excel,make_pdf,summary_cards
from test_core import HEADERS,rec,fx

class NegativeCommissionTests(unittest.TestCase):
    def test_negative_lines_remain_visible_with_zero_commission(self):
        with tempfile.TemporaryDirectory() as tmp:
            store=Store(Path(tmp)/'db.sqlite3')
            records=[rec('positive',margin='100',currency='HUF'),
                     rec('negative-huf',margin='-5000',currency='HUF'),
                     rec('negative-eur',margin='-246.21',currency='EUR'),
                     rec('zero',margin='0',currency='HUF')]
            p=build_plan(records,HEADERS,store,'2026-07','Peti','test',rate_provider=fx,snapshot=True)
            self.assertEqual(len(p['lines']),4)
            self.assertEqual(p['total'],10)
            self.assertEqual(p['current'],10)
            self.assertTrue(all(x['commission']==0 for x in p['lines'] if x['negative']))
            self.assertEqual(sum(x['negative'] for x in p['lines']),2)
            x=make_excel(p);wb=openpyxl.load_workbook(io.BytesIO(x),data_only=True)
            ws=wb['Jutalékelszámolás']
            for i,line in enumerate(p['lines'],9):
                if line['negative']:
                    self.assertLess(Decimal(line['original'][MARGIN]),0)
                    self.assertEqual(ws.cell(i,9).value,0)
                    self.assertEqual(ws.cell(i,9).font.color.rgb,'009C2222')
                    self.assertEqual(ws.cell(i,9).fill.fgColor.rgb,'00FFF0ED')
                    self.assertIn('Ft',ws.cell(i,9).number_format)
            pdf=make_pdf(p);reader=PdfReader(io.BytesIO(pdf))
            text=''.join(page.extract_text() for page in reader.pages)
            self.assertIn('-246,21 EUR',text)
            self.assertIn('-5 000,00 HUF',text)
            self.assertIn('0 Ft',text)
            self.assertIn('negatív árrésű tételek jutaléka 0 Ft',text)
            self.assertNotIn('csökkenti a jutalékot',text)
            self.assertIn('negatív árrésű tételek jutaléka 0 Ft',ws.cell(6,1).value)

if __name__=='__main__':unittest.main()
