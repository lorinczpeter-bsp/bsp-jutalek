"""Regression tests for month assignment across Streamlit reruns."""
import io,os,runpy,sys,tempfile,types,unittest,sqlite3
from pathlib import Path
from unittest.mock import patch
import openpyxl
from core import Store,origin,load_workbook
from test_core import HEADERS,rec,fx

class Rerun(Exception):pass
class State(dict):
    __getattr__=dict.get
    __setattr__=dict.__setitem__
class Surface:
    def __init__(self,app):self.app=app
    def __enter__(self):return self
    def __exit__(self,*args):pass
    def __getattr__(self,name):return getattr(self.app,name)
class App(types.ModuleType):
    def __init__(self,data):
        super().__init__('streamlit')
        self.session_state=State();self.secrets={};self.sidebar=Surface(self)
        self.data=data;self.submit=True;self.messages=[];self.options=[]
    def tabs(self,names):return [Surface(self) for _ in names]
    def columns(self,n):return [Surface(self) for _ in range(n)]
    def spinner(self,*args,**kwargs):return Surface(self)
    def form(self,*args,**kwargs):return Surface(self)
    def expander(self,label,*args,**kwargs):self.messages.append(label);return Surface(self)
    def file_uploader(self,*args,**kwargs):return None if kwargs.get('key')=='paid_report' else io.BytesIO(self.data)
    def selectbox(self,label,options,index=0,**kwargs):
        if label=='Ellenőrizendő fuvar':self.options.append(list(options))
        return options[index]
    def text_input(self,*args,**kwargs):return kwargs.get('value','')
    def number_input(self,*args,**kwargs):return kwargs.get('value',0)
    def form_submit_button(self,label,*args,**kwargs):return self.submit and label=='Hónap hozzárendelése'
    def rerun(self):raise Rerun()
    def stop(self):raise AssertionError('Unexpected stop')
    def success(self,message):self.messages.append(message)
    def error(self,message):self.messages.append(message)
    def __getattr__(self,name):return lambda *args,**kwargs:None

class MonthAssignmentTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.db=Path(self.tmp.name)/'state.sqlite3'
        wb=openpyxl.Workbook();ws=wb.active;ws.append(HEADERS)
        ws.append(list(rec('cross',dates='2026-09-01, 2026-10-01').original.values()))
        out=io.BytesIO();wb.save(out);self.data=out.getvalue()
        self.headers,self.records,_=load_workbook(self.data)
        Store(self.db).save_rate('2026-09',fx('2026-09'))
        self.ui=App(self.data)
    def tearDown(self):self.tmp.cleanup()
    def run_app(self):
        with patch.dict(sys.modules,{'streamlit':self.ui}),patch.dict(os.environ,{'JUTALEK_DB':str(self.db),'JUTALEK_PASSWORD':''}):
            runpy.run_path(str(Path(__file__).with_name('app.py')),run_name='__main__')
    def test_saved_with_blank_note_and_removed_after_rerun(self):
        with self.assertRaises(Rerun):self.run_app()
        self.assertEqual(origin(self.records[0],Store(self.db)),'2026-09')
        self.assertIn('cross',self.ui.session_state['decision_feedback'])
        self.ui.submit=False;self.ui.options=[];self.run_app()
        self.assertEqual(self.ui.options,[])
        self.assertTrue(any('Sikeres hozzárendelés: cross' in x for x in self.ui.messages))
        self.assertTrue(any('Rendezett hónap-hozzárendelések (1)' in x for x in self.ui.messages))
        self.assertNotIn('decision_feedback',self.ui.session_state)
    def test_database_failure_is_visible_and_remains_unresolved(self):
        with patch.object(Store,'save_decision',side_effect=sqlite3.OperationalError('database is locked')):
            self.run_app()
        self.assertIsNone(origin(self.records[0],Store(self.db)))
        self.assertTrue(any('Sikertelen hónap-hozzárendelés: database is locked' in x for x in self.ui.messages))
        self.assertEqual(self.ui.options,[['cross']])
    def test_two_assignments_remove_each_saved_freight(self):
        wb=openpyxl.load_workbook(io.BytesIO(self.data));ws=wb.active
        ws.append(list(rec('cross-2',dates='2026-09-03, 2026-10-02').original.values()))
        out=io.BytesIO();wb.save(out);self.ui.data=out.getvalue()
        with self.assertRaises(Rerun):self.run_app()
        with self.assertRaises(Rerun):self.run_app()
        self.assertEqual(self.ui.options,[['cross','cross-2'],['cross-2']])
        self.ui.submit=False;self.run_app()
        _,records,_=load_workbook(self.ui.data)
        self.assertTrue(all(origin(r,Store(self.db))=='2026-09' for r in records))

if __name__=='__main__':unittest.main()
